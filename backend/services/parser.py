import fitz
from docx import Document as DocxDocument
from pathlib import Path
from backend.schemas.document import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


class ParserService:

    @staticmethod
    def parse_txt(file_path: str, original_filename: str = None) -> Document:
        text = Path(file_path).read_text(
            encoding="utf-8",
            errors="ignore"
        )
        return Document(
            content=text,
            metadata={
                "source": file_path,
                "type": "txt"
            }
        )

    @staticmethod
    def parse_md(file_path: str, original_filename: str = None) -> Document:
        text = Path(file_path).read_text(
            encoding="utf-8",
            errors="ignore"
        )
        return Document(
            content=text,
            metadata={
                "source": file_path,
                "type": "md"
            }
        )

    @staticmethod
    def parse_docx(file_path: str, original_filename: str = None) -> Document:
        doc = DocxDocument(file_path)
        content_parts = []

        for element in doc.element.body:
            if element.tag.endswith('p'):
                p = Paragraph(element, doc)
                text = p.text.strip()
                if not text:
                    continue

                style_name = p.style.name if p.style else ""

                if style_name.startswith('Heading 1'):
                    content_parts.append(f"\n# {text}\n")
                elif style_name.startswith('Heading 2'):
                    content_parts.append(f"\n## {text}\n")
                elif style_name.startswith('Heading 3'):
                    content_parts.append(f"\n### {text}\n")
                else:
                    content_parts.append(text)

            elif element.tag.endswith('tbl'):
                table = Table(element, doc)
                md_table = []
                for row_idx, row in enumerate(table.rows):
                    row_cells = [cell.text.strip().replace("\n", " ") or "-"
                                 for cell in row.cells
                                 ]
                    md_table.append("| " + " | ".join(row_cells) + " | ")
                    if row_idx == 0:
                        md_table.append(
                            "| " + " | ".join(["---"] * len(row_cells)) + " |")
                if md_table:
                    content_parts.append("\n" + "\n".join(md_table) + "\n")

        return Document(
            content="\n\n".join(content_parts),
            metadata={"source": file_path,
                      "type": "docx"
                      },
        )

    @staticmethod
    def parse_pdf(file_path: str, original_filename: str = None) -> Document:
        text_parts = []
        with fitz.open(file_path) as pdf:
            for page_idx, page in enumerate(pdf):
                rect = page.rect
                header_margin = 50
                footer_margin = rect.height - 50

                blocks = page.get_text("blocks")
                page_text = []
                for b in blocks:
                    y0, y1 = b[1], b[3]
                    if y0 >= header_margin and y1 <= footer_margin:
                        content = b[4].strip()
                        if content:
                            page_text.append(content)
                if page_text:
                    text_parts.append("\n".join(page_text))
        return Document(
            content="\n\n".join(text_parts),
            metadata={"source": file_path, "type": "pdf"}
        )

    @staticmethod
    def parse(file_path: str) -> Document:
        suffix = Path(file_path).suffix.lower()
        if suffix == ".txt":
            return ParserService.parse_txt(file_path)
        elif suffix == ".md":
            return ParserService.parse_md(file_path)
        elif suffix == ".docx":
            return ParserService.parse_docx(file_path)
        elif suffix == ".pdf":
            return ParserService.parse_pdf(file_path)
        else:
            raise ValueError(f"Unsupported file type: {suffix}")
